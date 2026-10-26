#!/usr/bin/env python3
"""Seed Pure Maths and DSE M2 Analyse maps from math/M2/PManalyse.xlsx.

The workbook is the file on the Synology share at math/M2/PManalyse.xlsx
(it records its own folder as W:\\M2\\). It covers both syllabi:

- The upper question block is Paper 1. Calendar years before 2012 are
  Pure Maths. SP, PP, and 2012 onward are DSE M2.
- The lower question block is Paper 2. Only calendar years before 2012
  are exam columns, and they are Pure Maths. The 2015/2016 columns in
  that block are the M2 name list, not Paper 2 scripts.
- Pure Maths topic tokens are whatever the cells print. That includes
  0-15 from the bottom legend and extra tokens such as 16, 17, 0X, 0x,
  14X, and 15X. They are not cut down to the M2 range.
- M2 topic tokens are the separate 1-13 name list. The same integer is a
  different topic on each syllabus.

Topic tokens are stored as text. 0X and 0x stay distinct. This script
does not read or write the compulsory-core Analyse tables or analyse_m1_*.

The default run only parses and prints a report. It does not connect to
MySQL. See db/README.md.

Examples:
  python3 db/seed_analyse_pm.py --dry-run
  python3 db/seed_analyse_pm.py --check
  python3 db/seed_analyse_pm.py --emit-sql > /tmp/pm_analyse.sql
  M2_ANALYSE_XLSX=/Volumes/math/M2/PManalyse.xlsx python3 db/seed_analyse_pm.py --apply
"""

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import warnings
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
SCHEMA_PATH = SCRIPT_DIR / "analyse_pm_schema.sql"
FIXTURE_PATH = SCRIPT_DIR / "fixtures" / "M2" / "PManalyse.xlsx"

# Canonical file on the Synology share `math`. The workbook's own absPath
# is W:\M2\. Compulsory core stays at math/HKDSE/Analyse.xls and M1 stays
# at math/M1/analyse.xlsx. Neither is opened here.
M2_DIR = Path("math") / "M2"
CANONICAL_NAME = "PManalyse.xlsx"
FALLBACK_NAMES = ("PManalysis.xlsx", "pmanalysis.xlsx")
CORE_RELATIVE = Path("math") / "HKDSE" / "Analyse.xls"

YEAR_HEADER = re.compile(r"^(?:\d{2,4}|[A-Za-z]{1,6})$")
TOPIC_PART = re.compile(r"^(\d+)([Xx])?$")
TOPIC_CELL = re.compile(r"^\d+[Xx]?(?:\s*,\s*\d+[Xx]?)*$")
PAPER_CODES = {"SP", "PP"}

SYLLABUS_PM = "pm"
SYLLABUS_M2 = "m2"


@dataclass(frozen=True)
class Unit:
    syllabus: str
    code: str
    name: str | None


@dataclass(frozen=True)
class Question:
    sheet_name: str
    paper: int
    year_key: str
    year_label: str
    label_occurrence: int
    source_column: str
    column_index: int
    source_row: int
    question_no: int
    raw_cell: str
    unit_codes: tuple[str, ...]
    syllabus: str


@dataclass(frozen=True)
class Notice:
    level: str
    message: str


@dataclass(frozen=True)
class WorkbookMap:
    path: Path
    units: tuple[Unit, ...]
    questions: tuple[Question, ...]
    notices: tuple[Notice, ...]

    @property
    def errors(self) -> tuple[Notice, ...]:
        return tuple(notice for notice in self.notices if notice.level == "error")


def candidate_paths() -> list[Path]:
    """Search order after --xlsx and M2_ANALYSE_XLSX.

    Canonical PManalyse.xlsx wins over the pmanalysis fallbacks. The
    committed fixture is last.
    """
    home = Path.home()
    roots = [
        Path("/volume1") / M2_DIR,
        Path("/Volumes/math/M2"),
        Path("/Volumes/DiskStation") / M2_DIR,
        home / "SynologyDrive" / M2_DIR,
        home / "DiskStation" / M2_DIR,
        home / "Desktop" / M2_DIR,
        Path("D:/math/M2"),
        Path("W:/M2"),
        Path.cwd() / M2_DIR,
    ]
    paths = [root / CANONICAL_NAME for root in roots]
    paths.extend(root / name for root in roots for name in FALLBACK_NAMES)
    paths.append(FIXTURE_PATH)
    return paths


def resolve_xlsx(explicit: str | None = None, env: str | None = None) -> Path:
    if explicit:
        path = Path(explicit).expanduser()
        if not path.is_file():
            raise FileNotFoundError(f"Pure Maths / M2 workbook not found: {path}")
        return path
    if env:
        path = Path(env).expanduser()
        if not path.is_file():
            raise FileNotFoundError(f"M2_ANALYSE_XLSX does not exist: {path}")
        return path
    for path in candidate_paths():
        if path.is_file():
            return path
    searched = "\n".join(f"  - {path}" for path in candidate_paths())
    raise FileNotFoundError(
        "Workbook math/M2/PManalyse.xlsx was not found. Pass --xlsx or set "
        f"M2_ANALYSE_XLSX.\nLooked in:\n{searched}"
    )


def calendar_year(label: str) -> int | None:
    """Map a header to a calendar year.

    This sheet starts at 80. Two-digit years 80-99 are 1980-1999, and
    00-79 are 2000-2079, so 80 → 1980, 00 → 2000, and 12 → 2012.
    That century split is wider than the M1 workbook, which starts at 94.
    """
    if not re.fullmatch(r"\d{2,4}", label):
        return None
    year = int(label)
    if year < 100:
        year += 1900 if year >= 80 else 2000
    return year


def column_sequence_key(label: str, column_index: int) -> tuple[int, int, int]:
    """Order columns as pre-2012, then SP, PP, then 2012 onward.

    PManalyse.xlsx already lays the headers out in that order. The sort
    keeps the seed order if a later save shuffles the columns.
    """
    token = label.strip().upper()
    if token == "SP":
        return (1, 0, column_index)
    if token == "PP":
        return (1, 1, column_index)
    year = calendar_year(label.strip())
    if year is not None and year < 2012:
        return (0, year, column_index)
    if year is not None:
        return (2, year, column_index)
    return (3, 0, column_index)


def paper1_syllabus(label: str) -> str:
    """Paper 1 syllabus for one year column.

    Years before 2012 are Pure Maths, including the years that no longer
    cite 16, 17, or 0X. SP, PP, and every column from 2012 onward are M2.
    This sheet has a single column labelled 12, so there is no M1-style
    first-12 / second-12 split.
    """
    token = label.strip().upper()
    if token in PAPER_CODES:
        return SYLLABUS_M2
    year = calendar_year(label.strip())
    if year is None:
        raise ValueError(f"year label {label!r} has no Pure Maths / M2 syllabus rule")
    if year < 2012:
        return SYLLABUS_PM
    return SYLLABUS_M2


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ValueError(f"unexpected boolean cell {value!r}")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"unexpected fractional cell {value!r}")
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def is_topic_text(text: str) -> bool:
    return bool(text) and TOPIC_CELL.fullmatch(text) is not None


def is_name(value: object) -> bool:
    if value is None or isinstance(value, (int, float, bool)):
        return False
    text = str(value).strip()
    return bool(text) and not is_topic_text(text)


def canonical_code(part: str) -> str:
    match = TOPIC_PART.fullmatch(part)
    if match is None:
        raise ValueError(f"topic token {part!r} is not a workbook topic")
    suffix = match.group(2) or ""
    return str(int(match.group(1))) + suffix


def parse_topic_codes(value: object) -> tuple[str, ...]:
    text = cell_text(value)
    if not text:
        return ()
    if not is_topic_text(text):
        raise ValueError(f"topic cell {text!r} is not a topic token from the workbook")
    return tuple(canonical_code(part) for part in re.split(r"\s*,\s*", text))


def code_sort_key(code: str) -> tuple[int, int, str]:
    match = TOPIC_PART.fullmatch(code)
    if match is None:
        return (10**9, 1, code)
    suffix = match.group(2) or ""
    return (int(match.group(1)), 0 if suffix == "" else 1, suffix)


def question_number(value: object) -> int | None:
    if value is None:
        return None
    text = cell_text(value)
    if not text or not re.fullmatch(r"\d+", text):
        return None
    return int(text)


def load_pm_workbook(path: Path) -> WorkbookMap:
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Unknown extension is not supported and will be removed",
            category=UserWarning,
        )
        workbook = load_workbook(path, data_only=False, read_only=False)

    notices: list[Notice] = []
    questions: list[Question] = []
    pm_names: dict[str, str] = {}
    m2_names: dict[str, str] = {}

    try:
        for sheet in workbook.worksheets:
            parsed = _parse_sheet(sheet, notices)
            if parsed is None:
                continue
            sheet_pm, sheet_m2, sheet_questions = parsed
            _merge_names(pm_names, sheet_pm, "Pure Maths")
            _merge_names(m2_names, sheet_m2, "M2")
            questions.extend(sheet_questions)
    finally:
        workbook.close()

    if not m2_names:
        raise ValueError("M2 name list was not found in the workbook")
    units = _build_units(pm_names, m2_names, questions, notices)
    notices.extend(_m2_range_errors(questions, set(m2_names)))
    _append_layout_notes(notices, m2_names)
    return WorkbookMap(
        path=path,
        units=units,
        questions=tuple(questions),
        notices=tuple(notices),
    )


def _merge_names(existing: dict[str, str], incoming: dict[str, str], label: str) -> None:
    for code, name in incoming.items():
        previous = existing.get(code)
        if previous is None:
            existing[code] = name
        elif previous != name:
            raise ValueError(
                f"{label} topic {code} has conflicting names: {previous!r} vs {name!r}"
            )


def _build_units(
    pm_names: dict[str, str],
    m2_names: dict[str, str],
    questions: list[Question],
    notices: list[Notice],
) -> tuple[Unit, ...]:
    used_pm = {
        code
        for question in questions
        if question.syllabus == SYLLABUS_PM
        for code in question.unit_codes
    }
    pm_codes = sorted(set(pm_names) | used_pm, key=code_sort_key)
    unnamed = [code for code in pm_codes if code not in pm_names]
    if unnamed:
        notices.append(
            Notice(
                "info",
                "pm topic tokens with no legend name: " + ", ".join(unnamed),
            )
        )
    units = [
        Unit(SYLLABUS_PM, code, pm_names.get(code))
        for code in pm_codes
    ]
    for code in sorted(m2_names, key=code_sort_key):
        units.append(Unit(SYLLABUS_M2, code, m2_names[code]))
    return tuple(units)


def _m2_range_errors(questions: list[Question], allowed: set[str]) -> list[Notice]:
    notices: list[Notice] = []
    ordered = ", ".join(sorted(allowed, key=code_sort_key))
    for question in questions:
        if question.syllabus != SYLLABUS_M2:
            continue
        outside = [code for code in question.unit_codes if code not in allowed]
        if not outside:
            continue
        notices.append(
            Notice(
                "error",
                f"{question.sheet_name}!{question.source_column}{question.source_row} "
                f"paper {question.paper} {question.year_key} Q{question.question_no} "
                f"syllabus m2: topic {', '.join(outside)} is outside the M2 legend "
                f"({ordered})",
            )
        )
    return notices


def _append_layout_notes(notices: list[Notice], m2_names: dict[str, str]) -> None:
    notices.append(
        Notice(
            "info",
            "Paper 1 before 2012 is Pure Maths. SP, PP, and 2012 onward on "
            "Paper 1 are M2. The sheet has one column labelled 12, so that "
            "column is M2.",
        )
    )
    notices.append(
        Notice(
            "info",
            "Paper 2 exam columns are calendar years before 2012 only. "
            "Later columns in the lower block are not Paper 2 scripts.",
        )
    )
    notices.append(
        Notice(
            "info",
            "0X and 0x are stored as different topic tokens.",
        )
    )
    notices.append(
        Notice(
            "info",
            "M2 legend: "
            + ", ".join(
                f"{code} {m2_names[code]}"
                for code in sorted(m2_names, key=code_sort_key)
            ),
        )
    )


def _parse_sheet(
    sheet,
    notices: list[Notice],
) -> tuple[dict[str, str], dict[str, str], list[Question]] | None:
    header_row = _find_header_row(sheet)
    if header_row is None:
        notices.append(Notice("info", f"{sheet.title}: no year header, skipped"))
        return None

    year_columns = _year_columns(sheet, header_row)
    if not year_columns:
        notices.append(Notice("info", f"{sheet.title}: no year columns, skipped"))
        return None

    blocks = _question_blocks(sheet, header_row, year_columns)
    if len(blocks) != 2:
        spans = ["-".join(str(row) for row in (block[0], block[-1])) for block in blocks]
        raise ValueError(
            f"{sheet.title}: expected two question blocks (upper Paper 1, "
            f"lower Paper 2), found {len(blocks)} ({', '.join(spans) or 'none'})"
        )

    paper_rows = {1: blocks[0], 2: blocks[1]}
    legend_columns, m2_names = _m2_legend(sheet, paper_rows[2], year_columns)
    pm_names = _pm_legend(sheet, header_row, paper_rows, m2_names, year_columns)
    questions = _parse_questions(sheet, year_columns, paper_rows, legend_columns)
    _reject_unexpected_paper2_cells(sheet, year_columns, paper_rows[2], legend_columns)

    notices.append(
        Notice(
            "info",
            f"{sheet.title}: paper 1 rows {paper_rows[1][0]}-{paper_rows[1][-1]}, "
            f"paper 2 rows {paper_rows[2][0]}-{paper_rows[2][-1]}, "
            f"{len(questions)} questions, {len(pm_names)} Pure Maths legend names, "
            f"{len(m2_names)} M2 legend names",
        )
    )
    return pm_names, m2_names, questions


def _find_header_row(sheet) -> int | None:
    limit = min(sheet.max_row, 5)
    for row in range(1, limit + 1):
        labels = [
            cell_text(sheet.cell(row, column).value)
            for column in range(2, sheet.max_column + 1)
            if sheet.cell(row, column).value is not None
        ]
        if sum(1 for label in labels if YEAR_HEADER.fullmatch(label)) >= 3:
            return row
    return None


def _year_columns(sheet, header_row: int) -> list[tuple[int, str, str]]:
    columns: list[tuple[int, str, str]] = []
    for column in range(2, sheet.max_column + 1):
        value = sheet.cell(header_row, column).value
        if value is None or cell_text(value) == "":
            continue
        if not YEAR_HEADER.fullmatch(cell_text(value)):
            raise ValueError(
                f"{sheet.title}!{get_column_letter(column)}{header_row} "
                f"header {value!r} is not a year or paper code"
            )
        label = cell_text(value)
        token = label.upper()
        if re.fullmatch(r"[A-Za-z]{1,6}", label) and token not in PAPER_CODES:
            raise ValueError(
                f"{sheet.title}!{get_column_letter(column)}{header_row} "
                f"paper code {label!r} has no Pure Maths / M2 rule "
                "(this workbook uses SP and PP only)"
            )
        columns.append((column, get_column_letter(column), label))

    labels = [label for _, _, label in columns]
    repeated = sorted({label for label, count in Counter(labels).items() if count > 1})
    if repeated:
        raise ValueError(
            f"{sheet.title}: repeated year header {', '.join(repeated)}. "
            "PManalyse.xlsx has one column per label. Refusing to apply the "
            "M1 dual-2012 rule."
        )
    labelled_2012 = [label for label in labels if calendar_year(label) == 2012]
    if len(labelled_2012) > 1:
        raise ValueError(
            f"{sheet.title}: more than one column falls in 2012 ({', '.join(labelled_2012)}). "
            "Refusing to apply the M1 dual-2012 rule."
        )
    return columns


def _question_blocks(
    sheet,
    header_row: int,
    year_columns: list[tuple[int, str, str]],
) -> list[list[int]]:
    numbered = [
        row
        for row in range(header_row + 1, sheet.max_row + 1)
        if (number := question_number(sheet.cell(row, 1).value)) is not None and number >= 1
    ]
    spans: list[list[int]] = []
    current: list[int] = []
    previous: int | None = None
    for row in numbered:
        if previous is not None and row != previous + 1:
            if current:
                spans.append(current)
            current = []
        current.append(row)
        previous = row
    if current:
        spans.append(current)

    blocks: list[list[int]] = []
    for span in spans:
        if _span_has_topic(sheet, span, year_columns):
            blocks.append(span)
    return blocks


def _span_has_topic(sheet, rows: list[int], year_columns: list[tuple[int, str, str]]) -> bool:
    for row in rows:
        for column, _, _ in year_columns:
            if is_topic_text(cell_text(sheet.cell(row, column).value)):
                return True
    return False


def _m2_legend(
    sheet,
    paper2_rows: list[int],
    year_columns: list[tuple[int, str, str]],
) -> tuple[set[int], dict[str, str]]:
    """Find the name list embedded in the Paper 2 rows.

    In this workbook that pair is the columns headed 15 and 16: the left
    cells are 1..13 in order and the right cells are the M2 names. Those
    cells are not 2015/2016 Paper 2 answers. Paper 1 of those same columns
    is a real M2 script and is parsed separately.
    """
    found: list[tuple[int, int, dict[str, str]]] = []
    indexes = [column for column, _, _ in year_columns]
    located = {column: (letter, label) for column, letter, label in year_columns}
    for left, right in zip(indexes, indexes[1:]):
        if right != left + 1:
            continue
        names = _legend_pair(sheet, paper2_rows, left, right)
        if names is None:
            continue
        left_year = calendar_year(located[left][1])
        right_year = calendar_year(located[right][1])
        if left_year is None or right_year is None or left_year < 2012 or right_year < 2012:
            raise ValueError(
                f"{sheet.title}: M2 name list sits on "
                f"{located[left][0]}/{located[right][0]} "
                f"({located[left][1]!r}, {located[right][1]!r}), which is not a "
                "2012-or-later column pair. Refusing to drop those exam cells."
            )
        found.append((left, right, names))

    if len(found) != 1:
        raise ValueError(
            f"{sheet.title}: expected one M2 name list in the Paper 2 rows, found {len(found)}"
        )
    left, right, names = found[0]
    return {left, right}, names


def _legend_pair(sheet, rows: list[int], left: int, right: int) -> dict[str, str] | None:
    pairs: list[tuple[str, str]] = []
    for row in rows:
        left_value = sheet.cell(row, left).value
        right_value = sheet.cell(row, right).value
        left_text = cell_text(left_value) if left_value is not None else ""
        right_text = cell_text(right_value) if right_value is not None else ""
        if left_text == "" and right_text == "":
            continue
        if left_text == "" or right_text == "" or not is_name(right_value):
            return None
        try:
            codes = parse_topic_codes(left_value)
        except ValueError:
            return None
        if len(codes) != 1 or not codes[0].isdigit():
            return None
        pairs.append((codes[0], right_text))
    if len(pairs) < 3:
        return None
    expected = [str(number) for number in range(1, len(pairs) + 1)]
    if [code for code, _ in pairs] != expected:
        return None
    return {code: name for code, name in pairs}


def _pm_legend(
    sheet,
    header_row: int,
    paper_rows: dict[int, list[int]],
    m2_names: dict[str, str],
    year_columns: list[tuple[int, str, str]],
) -> dict[str, str]:
    """Bottom block: column A is the Pure Maths topic, column B is its name.

    Column I of those rows repeats the M2 names. Those cells are checked
    against the Paper 2 name list and are not attached to the Pure Maths codes.
    """
    question_rows = {row for rows in paper_rows.values() for row in rows}
    occupied = {column for column, _, _ in year_columns}
    names: dict[str, str] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        if row in question_rows:
            continue
        number = question_number(sheet.cell(row, 1).value)
        if number is None:
            if _row_has_values(sheet, row):
                raise ValueError(
                    f"{sheet.title}!A{row} is outside the question blocks and the "
                    "Pure Maths legend, but the row is not empty"
                )
            continue
        raw_name = sheet.cell(row, 2).value
        if not is_name(raw_name):
            raise ValueError(
                f"{sheet.title}!B{row} should be the Pure Maths name for topic {number}"
            )
        code = str(number)
        if code in names:
            raise ValueError(f"{sheet.title}!A{row} repeats Pure Maths topic {code}")
        names[code] = str(raw_name).strip()
        _check_legend_scratch(sheet, row, code, m2_names, occupied)
    if not names:
        raise ValueError(f"{sheet.title}: Pure Maths legend (column A/B below the papers) is missing")
    return names


def _row_has_values(sheet, row: int) -> bool:
    for column in range(1, sheet.max_column + 1):
        if sheet.cell(row, column).value is not None:
            return True
    return False


def _check_legend_scratch(
    sheet,
    row: int,
    code: str,
    m2_names: dict[str, str],
    occupied: set[int],
) -> None:
    for column in range(3, sheet.max_column + 1):
        value = sheet.cell(row, column).value
        if value is None or cell_text(value) == "":
            continue
        if column not in occupied:
            raise ValueError(
                f"{sheet.title}!{get_column_letter(column)}{row} is outside the year grid"
            )
        text = cell_text(value)
        expected = m2_names.get(code)
        if expected is None or text != expected:
            raise ValueError(
                f"{sheet.title}!{get_column_letter(column)}{row} is {text!r} on the "
                f"Pure Maths legend row for {code}. It does not match the M2 name list, "
                "so it is not treated as a Pure Maths name."
            )


def _parse_questions(
    sheet,
    year_columns: list[tuple[int, str, str]],
    paper_rows: dict[int, list[int]],
    legend_columns: set[int],
) -> list[Question]:
    questions: list[Question] = []
    ordered = sorted(year_columns, key=lambda item: column_sequence_key(item[2], item[0]))
    for paper in (1, 2):
        columns = ordered
        if paper == 2:
            columns = [
                item
                for item in ordered
                if item[0] not in legend_columns and _is_paper2_exam(item[2])
            ]
        for row in paper_rows[paper]:
            number = question_number(sheet.cell(row, 1).value)
            if number is None:
                raise ValueError(f"{sheet.title}!A{row} is in a question block but is not a question number")
            for column, letter, label in columns:
                value = sheet.cell(row, column).value
                if value is None or cell_text(value) == "":
                    continue
                try:
                    codes = parse_topic_codes(value)
                except ValueError as exc:
                    raise ValueError(f"{sheet.title}!{letter}{row}: {exc}") from exc
                syllabus = paper1_syllabus(label) if paper == 1 else SYLLABUS_PM
                questions.append(
                    Question(
                        sheet_name=sheet.title,
                        paper=paper,
                        year_key=label,
                        year_label=label,
                        label_occurrence=1,
                        source_column=letter,
                        column_index=column,
                        source_row=row,
                        question_no=number,
                        raw_cell=cell_text(value),
                        unit_codes=codes,
                        syllabus=syllabus,
                    )
                )
    return questions


def _is_paper2_exam(label: str) -> bool:
    year = calendar_year(label.strip())
    return year is not None and year < 2012


def _reject_unexpected_paper2_cells(
    sheet,
    year_columns: list[tuple[int, str, str]],
    paper2_rows: list[int],
    legend_columns: set[int],
) -> None:
    for column, letter, label in year_columns:
        if column in legend_columns or _is_paper2_exam(label):
            continue
        for row in paper2_rows:
            value = sheet.cell(row, column).value
            if value is None or cell_text(value) == "":
                continue
            raise ValueError(
                f"{sheet.title}!{letter}{row} is in the Paper 2 block under header "
                f"{label!r}. Paper 2 exam columns stop before 2012, and this cell "
                "is not the M2 name list."
            )


def render_sql(mapping: WorkbookMap) -> str:
    if mapping.errors:
        details = "\n".join(notice.message for notice in mapping.errors)
        raise ValueError(f"refusing to render SQL with parse errors:\n{details}")

    lines = [
        "-- Generated by db/seed_analyse_pm.py",
        f"-- Source: {mapping.path}",
        "-- Topic tokens are the workbook text. pm is Pure Maths. m2 is the 1-13 name list.",
        "-- Paper 1 is the upper block. Paper 2 is the lower block.",
        "SET NAMES utf8mb4;",
        "START TRANSACTION;",
        "DELETE FROM analyse_pm_sq_units;",
        "DELETE FROM analyse_pm_sq;",
        "DELETE FROM analyse_pm_units;",
    ]
    if mapping.units:
        lines.append("INSERT INTO analyse_pm_units (syllabus, unit_code, name) VALUES")
        lines.append(",\n".join(_unit_values(unit) for unit in mapping.units) + ";")
    if mapping.questions:
        lines.append(
            "INSERT INTO analyse_pm_sq "
            "(id, sheet_name, paper, syllabus, year_key, year_label, label_occurrence, "
            "source_column, question_no, raw_cell) VALUES"
        )
        question_rows = []
        for sq_id, question in enumerate(mapping.questions, start=1):
            question_rows.append(
                "("
                f"{sq_id}, {sql_str(question.sheet_name)}, {question.paper}, "
                f"{sql_str(question.syllabus)}, {sql_str(question.year_key)}, "
                f"{sql_str(question.year_label)}, {question.label_occurrence}, "
                f"{sql_str(question.source_column)}, {question.question_no}, "
                f"{sql_str(question.raw_cell)}"
                ")"
            )
        lines.append(",\n".join(question_rows) + ";")
        lines.append(
            "INSERT INTO analyse_pm_sq_units (sq_id, position, syllabus, unit_code) VALUES"
        )
        link_rows = []
        for sq_id, question in enumerate(mapping.questions, start=1):
            for position, code in enumerate(question.unit_codes, start=1):
                link_rows.append(
                    f"({sq_id}, {position}, {sql_str(question.syllabus)}, {sql_str(code)})"
                )
        lines.append(",\n".join(link_rows) + ";")
    lines.append("COMMIT;")
    lines.append("")
    return "\n".join(lines)


def _unit_values(unit: Unit) -> str:
    return f"({sql_str(unit.syllabus)}, {sql_str(unit.code)}, {sql_nullable(unit.name)})"


def sql_nullable(value: str | None) -> str:
    if value is None:
        return "NULL"
    return sql_str(value)


def sql_str(value: str) -> str:
    escaped = value.replace("\\", "\\\\").replace("'", "''")
    return f"'{escaped}'"


def render_report(mapping: WorkbookMap) -> str:
    link_count = sum(len(question.unit_codes) for question in mapping.questions)
    lines = [f"source: {mapping.path}"]
    for syllabus in (SYLLABUS_PM, SYLLABUS_M2):
        group = [question for question in mapping.questions if question.syllabus == syllabus]
        group_links = sum(len(question.unit_codes) for question in group)
        group_units = [unit.code for unit in mapping.units if unit.syllabus == syllabus]
        lines.append(
            f"syllabus {syllabus}: units {len(group_units)} "
            f"({', '.join(group_units)}), questions {len(group)}, topic links {group_links}"
        )
    for paper in (1, 2):
        keys: list[str] = []
        seen: set[str] = set()
        for question in mapping.questions:
            if question.paper != paper or question.year_key in seen:
                continue
            seen.add(question.year_key)
            keys.append(question.year_key)
        lines.append(f"paper {paper} year_key: " + ", ".join(keys))
    lines.append(f"questions: {len(mapping.questions)}")
    lines.append(f"topic links: {link_count}")
    for notice in mapping.notices:
        lines.append(f"{notice.level}: {notice.message}")
    lines.append("result: failed" if mapping.errors else "result: ok")
    return "\n".join(lines)


def apply_sql(sql: str) -> None:
    mysql = os.environ.get("MYSQL_BIN", "mysql")
    database = os.environ.get("MYSQL_DATABASE", "")
    if not database:
        raise SystemExit(
            "Refusing to apply: set MYSQL_DATABASE (and MYSQL_HOST, MYSQL_PORT, "
            "MYSQL_USER, MYSQL_PWD) on the machine that hosts alphaeqt. "
            "This script does not contact a database during --dry-run."
        )
    command = [
        mysql,
        "--batch",
        "--connect-timeout=10",
        "-h",
        os.environ.get("MYSQL_HOST", "127.0.0.1"),
        "-P",
        os.environ.get("MYSQL_PORT", "3306"),
        "-u",
        os.environ.get("MYSQL_USER", "root"),
    ]
    defaults = os.environ.get("MYSQL_DEFAULTS_FILE")
    if defaults:
        command.extend([f"--defaults-extra-file={defaults}"])
    command.append(database)
    completed = subprocess.run(command, input=sql, text=True, check=False)
    if completed.returncode != 0:
        raise SystemExit(f"mysql exited {completed.returncode}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Import Pure Maths and M2 Analyse topic tokens from "
            "math/M2/PManalyse.xlsx into analyse_pm_* tables. "
            "The default is a dry run and does not touch MySQL. "
            f"Compulsory core workbook {CORE_RELATIVE.as_posix()} is not read."
        )
    )
    parser.add_argument(
        "--xlsx",
        help=(
            "Path to PManalyse.xlsx. Overrides M2_ANALYSE_XLSX and the "
            "Synology search path. Canonical location: math/M2/PManalyse.xlsx "
            "(often /Volumes/math/M2/PManalyse.xlsx, D:\\math\\M2\\PManalyse.xlsx, "
            "or W:\\M2\\PManalyse.xlsx)."
        ),
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Parse and print a report. This is the default.",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help=(
            "Parse and exit 1 when an M2 cell uses a token outside the M2 "
            "name list, or when the sheet layout cannot be read. Pure Maths "
            "tokens are not capped at the M2 range."
        ),
    )
    parser.add_argument(
        "--emit-sql",
        action="store_true",
        help="Print SQL that replaces analyse_pm_* rows. Does not connect.",
    )
    parser.add_argument(
        "--with-schema",
        action="store_true",
        help="With --emit-sql or --apply, prepend db/analyse_pm_schema.sql.",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help=(
            "Apply schema and seed SQL through the local mysql client. "
            "Requires MYSQL_DATABASE. Do not point this at production from a "
            "machine that does not already have the alphaeqt credentials."
        ),
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    selected = sum(bool(flag) for flag in (args.check, args.emit_sql, args.apply))
    if args.dry_run and selected:
        print("--dry-run cannot be combined with --check, --emit-sql, or --apply", file=sys.stderr)
        return 2
    if selected > 1:
        print("use only one of --check, --emit-sql, --apply", file=sys.stderr)
        return 2

    try:
        path = resolve_xlsx(args.xlsx, os.environ.get("M2_ANALYSE_XLSX"))
        mapping = load_pm_workbook(path)
    except (FileNotFoundError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    if mapping.errors:
        print(render_report(mapping), file=sys.stderr)
        return 1

    if args.emit_sql or args.apply:
        sql = render_sql(mapping)
        if args.with_schema or args.apply:
            sql = SCHEMA_PATH.read_text(encoding="utf-8") + "\n" + sql
        if args.emit_sql:
            sys.stdout.write(sql)
            return 0
        apply_sql(sql)
        print(render_report(mapping))
        return 0

    print(render_report(mapping))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
