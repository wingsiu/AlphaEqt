#!/usr/bin/env python3
"""Verify M1 Analyse parsing keeps the topic numbers from analyse.xlsx."""

from __future__ import annotations

import importlib.util
import re
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
FIXTURE = ROOT / "fixtures" / "M1" / "analyse.xlsx"


def load_seeder():
    spec = importlib.util.spec_from_file_location("seed_analyse_m1", ROOT / "seed_analyse_m1.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("could not load seed_analyse_m1.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class SeedAnalyseM1Test(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.mod = load_seeder()
        cls.mapping = cls.mod.load_m1_workbook(FIXTURE)

    def test_fixture_exists(self):
        self.assertTrue(FIXTURE.is_file())

    def test_topic_numbers_match_excel_codes(self):
        ms = [unit.code for unit in self.mapping.units if unit.syllabus == "ms"]
        dse = [unit.code for unit in self.mapping.units if unit.syllabus == "dse"]
        self.assertEqual(ms, list(range(1, 17)))
        self.assertEqual(dse, list(range(1, 13)))
        by_key = {(unit.syllabus, unit.code): unit for unit in self.mapping.units}
        self.assertEqual(by_key[("ms", 1)].short_name, "binomial")
        self.assertEqual(by_key[("dse", 1)].name, "Binomial")
        self.assertEqual(by_key[("ms", 3)].short_name, "limit")
        self.assertEqual(by_key[("dse", 5)].name, "Application of Differentiation(1)")
        self.assertIsNone(by_key[("ms", 13)].short_name)
        self.assertIsNone(by_key[("ms", 13)].name)
        self.assertIsNone(by_key[("ms", 15)].name)
        self.assertIsNone(by_key[("ms", 16)].short_name)
        self.assertNotIn(("dse", 13), by_key)
        self.assertNotIn(("dse", 16), by_key)

    def test_grid_counts_and_cell_order(self):
        self.assertEqual(len(self.mapping.questions), 448)
        links = sum(len(question.unit_codes) for question in self.mapping.questions)
        self.assertEqual(links, 498)
        self.assertEqual(self.mapping.errors, ())

        def find(year_key: str, question_no: int):
            matches = [
                question
                for question in self.mapping.questions
                if question.year_key == year_key and question.question_no == question_no
            ]
            self.assertEqual(len(matches), 1)
            return matches[0]

        self.assertEqual(find("94", 1).unit_codes, (12,))
        self.assertEqual(find("94", 1).syllabus, "ms")
        self.assertEqual(find("94", 1).source_column, "B")
        self.assertEqual(find("00", 1).year_label, "00")
        self.assertEqual(find("00", 1).syllabus, "ms")
        self.assertEqual(find("00", 1).unit_codes, (6,))
        self.assertEqual(find("09", 8).raw_cell, "4, 7")
        self.assertEqual(find("09", 8).unit_codes, (4, 7))
        self.assertEqual(find("11", 10).unit_codes, (14,))
        self.assertEqual(find("11", 10).syllabus, "ms")
        self.assertEqual(find("25", 1).unit_codes, (10, 12))
        self.assertEqual(find("25", 1).syllabus, "dse")
        self.assertEqual(find("SP", 14).unit_codes, (12,))
        self.assertEqual(find("SP", 14).syllabus, "dse")
        self.assertEqual(find("PP", 1).syllabus, "dse")
        self.assertEqual(find("12#T", 2).unit_codes, (4,))
        self.assertEqual(find("12#T", 11).unit_codes, (13,))
        self.assertEqual(find("12#T", 11).syllabus, "ms")
        self.assertEqual(find("12#T", 2).syllabus, "ms")
        self.assertEqual(find("12#U", 2).unit_codes, (8,))
        self.assertEqual(find("12#U", 2).syllabus, "dse")
        self.assertEqual(find("12#T", 2).year_label, "12")
        self.assertEqual(find("12#T", 2).label_occurrence, 1)
        self.assertEqual(find("12#U", 2).label_occurrence, 2)
        self.assertEqual(find("26", 1).unit_codes, (10,))
        self.assertEqual(find("26", 1).syllabus, "dse")

    def test_year_key_order_is_sp_then_pp_then_2012(self):
        keys: list[str] = []
        for question in self.mapping.questions:
            if question.year_key not in keys:
                keys.append(question.year_key)
        self.assertLess(keys.index("11"), keys.index("SP"))
        self.assertLess(keys.index("SP"), keys.index("PP"))
        self.assertLess(keys.index("PP"), keys.index("12#T"))
        self.assertLess(keys.index("12#T"), keys.index("12#U"))
        self.assertLess(keys.index("12#U"), keys.index("13"))
        self.assertEqual(
            keys[keys.index("11"): keys.index("13") + 1],
            ["11", "SP", "PP", "12#T", "12#U", "13"],
        )

    def test_every_grid_cell_uses_the_excel_topic_numbers(self):
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Unknown extension is not supported and will be removed",
                category=UserWarning,
            )
            workbook = load_workbook(FIXTURE, data_only=False)
        sheet = workbook["Sheet1"]
        try:
            headers = {}
            counts = {}
            for column in range(2, sheet.max_column + 1):
                value = sheet.cell(1, column).value
                if value is None:
                    continue
                label = str(value).strip() if not isinstance(value, int) else str(value)
                counts[label] = counts.get(label, 0) + 1
                headers[column] = (get_column_letter(column), label, counts[label])
            seen = 0
            for row in range(2, 16):
                question_no = int(sheet.cell(row, 1).value)
                for column, (letter, label, occurrence) in headers.items():
                    value = sheet.cell(row, column).value
                    if value is None:
                        continue
                    raw = str(value).strip() if not isinstance(value, int) else str(value)
                    expected = tuple(int(part) for part in re.split(r"\s*,\s*", raw))
                    year_key = label if counts[label] == 1 else f"{label}#{letter}"
                    match = next(
                        question
                        for question in self.mapping.questions
                        if question.year_key == year_key and question.question_no == question_no
                    )
                    self.assertEqual(match.unit_codes, expected)
                    self.assertEqual(match.raw_cell, raw)
                    self.assertEqual(match.label_occurrence, occurrence)
                    self.assertEqual(match.source_column, letter)
                    self.assertEqual(
                        match.syllabus,
                        self.mod.syllabus_for(label, occurrence),
                    )
                    seen += 1
            self.assertEqual(seen, len(self.mapping.questions))
        finally:
            workbook.close()

    def test_empty_sheets_are_skipped(self):
        messages = [notice.message for notice in self.mapping.notices]
        self.assertTrue(any(message.startswith("Sheet2:") for message in messages))
        self.assertTrue(any(message.startswith("Sheet3:") for message in messages))
        sheets = {question.sheet_name for question in self.mapping.questions}
        self.assertEqual(sheets, {"Sheet1"})

    def test_sql_uses_excel_codes_and_skips_core_tables(self):
        sql = self.mod.render_sql(self.mapping)
        self.assertIn("('ms', 1, 'binomial', 'Binomial')", sql)
        self.assertIn("('dse', 1, 'binomial', 'Binomial')", sql)
        self.assertIn("('ms', 13, NULL, NULL)", sql)
        self.assertIn("('ms', 15, NULL, NULL)", sql)
        self.assertNotIn("('dse', 13,", sql)
        self.assertNotIn("('dse', 16,", sql)
        self.assertIn("'12#U'", sql)
        self.assertIn("'00'", sql)
        self.assertNotIn("2000", sql)
        self.assertIn("'12#T', '12', 1, 'T', 11, '13', 'ms'", sql)
        self.assertNotRegex(sql, r"'dse', 13\)")
        for name in (
            "analyse_mc",
            "analyse_mc_units",
            "analyse_sq",
            "analyse_sq_units",
            "analyse_units",
            "analyse_unit_qtypes",
        ):
            self.assertNotRegex(sql, rf"\b{name}\b")

    def test_schema_does_not_create_core_tables(self):
        schema = (ROOT / "analyse_m1_schema.sql").read_text(encoding="utf-8")
        self.assertIn("CREATE TABLE IF NOT EXISTS analyse_m1_units", schema)
        self.assertIn("CREATE TABLE IF NOT EXISTS analyse_m1_sq", schema)
        self.assertIn("CREATE TABLE IF NOT EXISTS analyse_m1_sq_units", schema)
        for name in (
            "analyse_mc",
            "analyse_mc_units",
            "analyse_sq",
            "analyse_sq_units",
            "analyse_units",
            "analyse_unit_qtypes",
        ):
            self.assertNotIn(f"CREATE TABLE IF NOT EXISTS {name}", schema)

    def test_first_2012_column_is_ms_and_the_second_is_dse(self):
        self.assertEqual(self.mod.syllabus_for("12", 1), "ms")
        self.assertEqual(self.mod.syllabus_for("12", 2), "dse")
        self.assertEqual(self.mod.syllabus_for("2012", 1), "ms")
        self.assertEqual(self.mod.syllabus_for("2012", 2), "dse")
        self.assertEqual(self.mod.syllabus_for("13", 1), "dse")
        self.assertEqual(self.mod.syllabus_for("SP", 1), "dse")
        self.assertEqual(self.mod.syllabus_for("PP", 1), "dse")

    def test_out_of_range_topics_fail_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "analyse.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Sheet1"
            sheet["A2"] = 1
            sheet["B1"] = "11"
            sheet["C1"] = 16
            sheet["D1"] = "SP"
            sheet["E1"] = "ON"
            sheet["B2"] = 16
            sheet["C2"] = 13
            sheet["D2"] = 13
            sheet["E2"] = 12
            book.save(path)
            mapping = self.mod.load_m1_workbook(path)
        keys: list[str] = []
        for question in mapping.questions:
            if question.year_key not in keys:
                keys.append(question.year_key)
        self.assertEqual(keys, ["11", "SP", "ON", "16"])
        by_year = {question.year_key: question for question in mapping.questions}
        self.assertEqual(by_year["11"].syllabus, "ms")
        self.assertEqual(by_year["11"].unit_codes, (16,))
        self.assertEqual(by_year["16"].syllabus, "dse")
        self.assertEqual(by_year["SP"].syllabus, "dse")
        self.assertEqual(by_year["ON"].syllabus, "dse")
        messages = [notice.message for notice in mapping.errors]
        self.assertTrue(any("16 Q1" in message and "topic 13" in message for message in messages))
        self.assertTrue(any("SP Q1" in message and "topic 13" in message for message in messages))
        self.assertFalse(any(message.startswith("Sheet1!") and " ON " in message for message in messages))
        self.assertFalse(any(" 11 Q1" in message for message in messages))
        dse_codes = [unit.code for unit in mapping.units if unit.syllabus == "dse"]
        self.assertEqual(dse_codes, list(range(1, 13)))

    def test_resolve_prefers_explicit_path(self):
        resolved = self.mod.resolve_xlsx(str(FIXTURE), None)
        self.assertEqual(resolved, FIXTURE)


if __name__ == "__main__":
    unittest.main()
