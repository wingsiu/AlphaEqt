#!/usr/bin/env python3
"""Verify Pure Maths and M2 Analyse parsing keeps the tokens in PManalyse.xlsx."""

from __future__ import annotations

import contextlib
import importlib.util
import io
import re
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
FIXTURE = ROOT / "fixtures" / "M2" / "PManalyse.xlsx"


def load_seeder():
    spec = importlib.util.spec_from_file_location("seed_analyse_pm", ROOT / "seed_analyse_pm.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("could not load seed_analyse_pm.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def raw_of(value: object) -> str:
    if isinstance(value, bool):
        raise AssertionError(value)
    if isinstance(value, float):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def tokens_of(raw: str) -> tuple[str, ...]:
    codes = []
    for part in re.split(r"\s*,\s*", raw):
        match = re.fullmatch(r"(\d+)([Xx])?", part)
        if match is None:
            raise AssertionError(part)
        codes.append(str(int(match.group(1))) + (match.group(2) or ""))
    return tuple(codes)


def calendar_year(label: str) -> int | None:
    if not re.fullmatch(r"\d{2,4}", label):
        return None
    year = int(label)
    if year < 100:
        year += 1900 if year >= 80 else 2000
    return year


class SeedAnalysePmTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.mod = load_seeder()
        cls.mapping = cls.mod.load_pm_workbook(FIXTURE)

    def test_fixture_exists(self):
        self.assertTrue(FIXTURE.is_file())

    def test_check_exits_zero(self):
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(self.mod.main(["--check", "--xlsx", str(FIXTURE)]), 0)
        self.assertEqual(self.mapping.errors, ())

    def test_topic_catalogues_stay_separate(self):
        pm = {unit.code: unit.name for unit in self.mapping.units if unit.syllabus == "pm"}
        m2 = {unit.code: unit.name for unit in self.mapping.units if unit.syllabus == "m2"}
        self.assertEqual(
            list(pm),
            [
                "0", "0X", "0x", "1", "2", "3", "4", "5", "6", "7", "8", "9",
                "10", "11", "12", "13", "14", "14X", "15", "15X", "16", "17",
            ],
        )
        self.assertEqual(list(m2), [str(code) for code in range(1, 14)])
        self.assertEqual(pm["0"], "Sets Language")
        self.assertEqual(pm["3"], "Polynomial")
        self.assertEqual(pm["13"], "integration")
        self.assertEqual(pm["15"], "analytical geometry")
        self.assertEqual(m2["1"], "MI")
        self.assertEqual(m2["2"], "Bionomial")
        self.assertEqual(m2["3"], "Trigonometry")
        self.assertEqual(m2["13"], "Vectors")
        for code in ("0X", "0x", "14X", "15X", "16", "17"):
            self.assertIsNone(pm[code])
        self.assertNotIn("16", m2)
        self.assertNotIn("0X", m2)

    def test_counts_and_anchor_cells(self):
        self.assertEqual(len(self.mapping.questions), 937)
        self.assertEqual(sum(len(q.unit_codes) for q in self.mapping.questions), 1045)

        def group(paper: int, syllabus: str):
            rows = [
                question
                for question in self.mapping.questions
                if question.paper == paper and question.syllabus == syllabus
            ]
            links = sum(len(question.unit_codes) for question in rows)
            return len(rows), links

        self.assertEqual(group(1, "pm"), (360, 397))
        self.assertEqual(group(1, "m2"), (215, 236))
        self.assertEqual(group(2, "pm"), (362, 412))
        self.assertEqual(group(2, "m2"), (0, 0))

        def find(paper: int, year_key: str, question_no: int):
            matches = [
                question
                for question in self.mapping.questions
                if question.paper == paper
                and question.year_key == year_key
                and question.question_no == question_no
            ]
            self.assertEqual(len(matches), 1)
            return matches[0]

        first = find(1, "80", 1)
        self.assertEqual(first.unit_codes, ("16", "6"))
        self.assertEqual(first.raw_cell, "16,6")
        self.assertEqual(first.syllabus, "pm")
        self.assertEqual(first.source_column, "B")
        self.assertEqual(first.label_occurrence, 1)
        self.assertEqual(find(1, "80", 2).unit_codes, ("0X",))
        self.assertEqual(find(1, "89", 7).unit_codes, ("0x",))
        self.assertEqual(find(1, "89", 7).raw_cell, "0x")
        self.assertEqual(find(2, "80", 1).unit_codes, ("15",))
        self.assertEqual(find(2, "80", 1).syllabus, "pm")
        self.assertEqual(find(2, "80", 1).source_row, 18)
        self.assertEqual(find(2, "80", 4).unit_codes, ("12", "15X"))
        self.assertEqual(find(2, "81", 4).unit_codes, ("3", "2", "11"))
        self.assertEqual(find(1, "11", 1).syllabus, "pm")
        self.assertEqual(find(1, "11", 1).unit_codes, ("1",))
        self.assertEqual(find(1, "SP", 1).syllabus, "m2")
        self.assertEqual(find(1, "SP", 1).unit_codes, ("4",))
        self.assertEqual(find(1, "PP", 1).syllabus, "m2")
        self.assertEqual(find(1, "12", 1).syllabus, "m2")
        self.assertEqual(find(1, "12", 1).unit_codes, ("4",))
        self.assertEqual(find(1, "12", 1).year_label, "12")
        spaced = find(1, "19", 4)
        self.assertEqual(spaced.raw_cell, "8, 10")
        self.assertEqual(spaced.unit_codes, ("8", "10"))
        self.assertEqual(find(1, "17", 9).raw_cell, "7, 10")
        self.assertEqual(find(1, "00", 1).year_label, "00")
        self.assertEqual(find(1, "00", 1).syllabus, "pm")
        self.assertEqual(find(1, "15", 1).syllabus, "m2")
        self.assertEqual(find(1, "16", 1).syllabus, "m2")

        paper2_years = {question.year_key for question in self.mapping.questions if question.paper == 2}
        self.assertTrue(paper2_years.isdisjoint({"SP", "PP", "12", "15", "16", "26"}))

    def test_year_order_is_pre_2012_then_sp_pp_then_2012(self):
        keys: list[str] = []
        for question in self.mapping.questions:
            if question.paper == 1 and question.year_key not in keys:
                keys.append(question.year_key)
        self.assertEqual(keys[0], "80")
        self.assertEqual(keys[-1], "26")
        self.assertEqual(
            keys[keys.index("11"): keys.index("13") + 1],
            ["11", "SP", "PP", "12", "13"],
        )
        paper2: list[str] = []
        for question in self.mapping.questions:
            if question.paper == 2 and question.year_key not in paper2:
                paper2.append(question.year_key)
        self.assertEqual(paper2[0], "80")
        self.assertEqual(paper2[-1], "11")
        self.assertNotIn("SP", paper2)

    def test_every_grid_cell_uses_the_excel_topic_tokens(self):
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Unknown extension is not supported and will be removed",
                category=UserWarning,
            )
            workbook = load_workbook(FIXTURE, data_only=False)
        sheet = workbook["Sheet1"]
        try:
            headers = []
            for column in range(2, sheet.max_column + 1):
                value = sheet.cell(1, column).value
                if value is None:
                    continue
                headers.append((column, get_column_letter(column), raw_of(value)))
            seen = 0
            for row in range(2, 16):
                question_no = int(sheet.cell(row, 1).value)
                seen += self._assert_row(sheet, row, question_no, 1, headers, skip=set())
            legend = self._paper2_legend_columns(sheet, headers)
            self.assertEqual(
                [label for _, _, label in headers if _column_of(headers, label) in legend],
                ["15", "16"],
            )
            for row in range(18, 32):
                question_no = int(sheet.cell(row, 1).value)
                seen += self._assert_row(sheet, row, question_no, 2, headers, skip=legend)
            self.assertEqual(seen, len(self.mapping.questions))
        finally:
            workbook.close()

    def _assert_row(self, sheet, row, question_no, paper, headers, skip) -> int:
        seen = 0
        for column, letter, label in headers:
            if column in skip:
                continue
            if paper == 2:
                year = calendar_year(label)
                if year is None or year >= 2012:
                    value = sheet.cell(row, column).value
                    self.assertTrue(value is None or raw_of(value) == "")
                    continue
            value = sheet.cell(row, column).value
            if value is None or raw_of(value) == "":
                continue
            raw = raw_of(value)
            match = next(
                question
                for question in self.mapping.questions
                if question.paper == paper
                and question.year_key == label
                and question.question_no == question_no
            )
            self.assertEqual(match.unit_codes, tokens_of(raw))
            self.assertEqual(match.raw_cell, raw)
            self.assertEqual(match.source_column, letter)
            self.assertEqual(match.label_occurrence, 1)
            if paper == 1:
                year = calendar_year(label)
                expected = "m2" if label in {"SP", "PP"} or (year is not None and year >= 2012) else "pm"
            else:
                expected = "pm"
            self.assertEqual(match.syllabus, expected)
            seen += 1
        return seen

    def _paper2_legend_columns(self, sheet, headers) -> set[int]:
        indexes = [column for column, _, _ in headers]
        found = []
        for left, right in zip(indexes, indexes[1:]):
            if right != left + 1:
                continue
            codes = []
            names = []
            for row in range(18, 32):
                left_value = sheet.cell(row, left).value
                right_value = sheet.cell(row, right).value
                if left_value is None and right_value is None:
                    continue
                codes.append(raw_of(left_value))
                names.append(raw_of(right_value))
            if codes == [str(number) for number in range(1, len(codes) + 1)] and len(codes) >= 3:
                found.append({left, right})
                self.assertEqual(names[0], "MI")
                self.assertEqual(names[1], "Bionomial")
        self.assertEqual(len(found), 1)
        return found[0]

    def test_century_split_starts_at_1980(self):
        self.assertEqual(self.mod.calendar_year("80"), 1980)
        self.assertEqual(self.mod.calendar_year("99"), 1999)
        self.assertEqual(self.mod.calendar_year("00"), 2000)
        self.assertEqual(self.mod.calendar_year("12"), 2012)
        self.assertEqual(self.mod.paper1_syllabus("80"), "pm")
        self.assertEqual(self.mod.paper1_syllabus("11"), "pm")
        self.assertEqual(self.mod.paper1_syllabus("SP"), "m2")
        self.assertEqual(self.mod.paper1_syllabus("PP"), "m2")
        self.assertEqual(self.mod.paper1_syllabus("12"), "m2")
        self.assertEqual(self.mod.paper1_syllabus("26"), "m2")

    def test_sql_uses_excel_tokens_and_skips_other_analyse_tables(self):
        sql = self.mod.render_sql(self.mapping)
        self.assertIn("('pm', '0', 'Sets Language')", sql)
        self.assertIn("('pm', '3', 'Polynomial')", sql)
        self.assertIn("('pm', '13', 'integration')", sql)
        self.assertIn("('pm', '0X', NULL)", sql)
        self.assertIn("('pm', '0x', NULL)", sql)
        self.assertIn("('pm', '14X', NULL)", sql)
        self.assertIn("('pm', '15X', NULL)", sql)
        self.assertIn("('pm', '16', NULL)", sql)
        self.assertIn("('m2', '2', 'Bionomial')", sql)
        self.assertIn("('m2', '3', 'Trigonometry')", sql)
        self.assertIn("('m2', '13', 'Vectors')", sql)
        self.assertIn("'00'", sql)
        self.assertNotIn("2000", sql)
        self.assertNotIn("analyse_m1_", sql)
        self.assertNotIn("analyse_m2_", sql)
        for name in (
            "analyse_mc",
            "analyse_mc_units",
            "analyse_sq",
            "analyse_sq_units",
            "analyse_units",
            "analyse_unit_qtypes",
        ):
            self.assertNotRegex(sql, rf"\b{name}\b")
        self.assertIn("DELETE FROM analyse_pm_sq_units;", sql)
        self.assertIn("DELETE FROM analyse_pm_sq;", sql)
        self.assertIn("DELETE FROM analyse_pm_units;", sql)

    def test_schema_does_not_create_core_or_m1_tables(self):
        schema = (ROOT / "analyse_pm_schema.sql").read_text(encoding="utf-8")
        self.assertIn("CREATE TABLE IF NOT EXISTS analyse_pm_units", schema)
        self.assertIn("CREATE TABLE IF NOT EXISTS analyse_pm_sq", schema)
        self.assertIn("CREATE TABLE IF NOT EXISTS analyse_pm_sq_units", schema)
        self.assertNotIn("CREATE TABLE IF NOT EXISTS analyse_m2_units", schema)
        self.assertNotIn("CREATE TABLE IF NOT EXISTS analyse_m1_units", schema)
        for name in (
            "analyse_mc",
            "analyse_sq",
            "analyse_units",
        ):
            self.assertNotIn(f"CREATE TABLE IF NOT EXISTS {name}", schema)

    def test_repeated_2012_header_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "PManalyse.xlsx"
            book = Workbook()
            sheet = book.active
            sheet["B1"] = "12"
            sheet["C1"] = "12"
            sheet["D1"] = "13"
            book.save(path)
            with self.assertRaises(ValueError) as raised:
                self.mod.load_pm_workbook(path)
        self.assertIn("dual-2012", str(raised.exception))

    def test_distinct_2012_labels_are_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "PManalyse.xlsx"
            book = Workbook()
            sheet = book.active
            sheet["B1"] = "12"
            sheet["C1"] = 2012
            sheet["D1"] = "13"
            book.save(path)
            with self.assertRaises(ValueError) as raised:
                self.mod.load_pm_workbook(path)
        self.assertIn("2012", str(raised.exception))

    def test_unknown_paper_code_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "PManalyse.xlsx"
            book = Workbook()
            sheet = book.active
            sheet["B1"] = "80"
            sheet["C1"] = "SP"
            sheet["D1"] = "ON"
            book.save(path)
            with self.assertRaises(ValueError) as raised:
                self.mod.load_pm_workbook(path)
        self.assertIn("ON", str(raised.exception))

    def test_m2_token_outside_the_legend_fails_and_pure_maths_does_not(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "PManalyse.xlsx"
            _write_minimal_workbook(Path(tmp) / "ok.xlsx")
            ok = self.mod.load_pm_workbook(Path(tmp) / "ok.xlsx")
            self.assertEqual(ok.errors, ())
            by_key = {(q.paper, q.year_key, q.question_no): q for q in ok.questions}
            self.assertEqual(by_key[(1, "80", 1)].unit_codes, ("17",))
            self.assertEqual(by_key[(1, "80", 1)].syllabus, "pm")
            self.assertEqual(by_key[(2, "80", 2)].unit_codes, ("14X",))
            self.assertIsNone(next(unit for unit in ok.units if unit.code == "17").name)
            paper1_keys: list[str] = []
            for question in ok.questions:
                if question.paper == 1 and question.year_key not in paper1_keys:
                    paper1_keys.append(question.year_key)
            self.assertEqual(paper1_keys, ["80", "SP", "12", "15", "16"])

            _write_minimal_workbook(path, m2_token="16")
            mapping = self.mod.load_pm_workbook(path)
        messages = [notice.message for notice in mapping.errors]
        self.assertTrue(any("SP Q1" in message and "topic 16" in message for message in messages))
        self.assertFalse(any(" 80 Q1" in message for message in messages))
        with self.assertRaises(ValueError):
            self.mod.render_sql(mapping)

    def test_resolve_prefers_explicit_path(self):
        resolved = self.mod.resolve_xlsx(str(FIXTURE), None)
        self.assertEqual(resolved, FIXTURE)


def _column_of(headers, label: str) -> int:
    return next(column for column, _, header in headers if header == label)


def _write_minimal_workbook(path: Path, m2_token: str = "1") -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet["B1"] = "80"
    sheet["C1"] = "SP"
    sheet["D1"] = "12"
    sheet["E1"] = 15
    sheet["F1"] = 16
    sheet["A2"] = 1
    sheet["B2"] = 17
    sheet["C2"] = m2_token
    sheet["D2"] = 2
    sheet["E2"] = 1
    sheet["F2"] = 2
    sheet["A3"] = 2
    sheet["B3"] = "0X"
    sheet["C3"] = 1
    sheet["D3"] = 2
    sheet["E3"] = 3
    sheet["F3"] = 1
    sheet["A5"] = 1
    sheet["B5"] = 15
    sheet["E5"] = 1
    sheet["F5"] = "MI "
    sheet["A6"] = 2
    sheet["B6"] = "14X"
    sheet["E6"] = 2
    sheet["F6"] = "Bionomial"
    sheet["A7"] = 3
    sheet["B7"] = 16
    sheet["E7"] = 3
    sheet["F7"] = "Trigonometry"
    sheet["A9"] = 0
    sheet["B9"] = "Sets Language"
    sheet["A10"] = 1
    sheet["B10"] = "Mathematical Induction"
    sheet["A11"] = 15
    sheet["B11"] = "analytical geometry"
    book.save(path)


if __name__ == "__main__":
    unittest.main()
