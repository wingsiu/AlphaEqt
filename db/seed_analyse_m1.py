#!/usr/bin/env python3
"""Seed HKDSE M1 Analyse maps from math/M1/analyse.xlsx.

Topic and unit identifiers are the topic numbers printed in the workbook
(legend column A and the year x question cells). Years before 2012, and the
first column labelled 12, use topics 1-16. The second column labelled 12,
later years, SP, PP, and ON use topics 1-12.
This script does not
renumber them and does not read or write the compulsory-core tables
(analyse_mc, analyse_sq, analyse_units, and the related junction tables).

The default run only parses and prints a report. It does not connect to
MySQL. Apply the result on a machine that can see the DiskStation share
and the alphaeqt database. See db/README.md.

Examples:
  python3 db/seed_analyse_m1.py --dry-run
  python3 db/seed_analyse_m1.py --check
  python3 db/seed_analyse_m1.py --emit-sql > /tmp/m1_analyse.sql
  M1_ANALYSE_XLSX=/Volumes/math/M1/analyse.xlsx python3 db/seed_analyse_m1.py --apply
"""

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import warnings
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
SCHEMA_PATH = SCRIPT_DIR / "analyse_m1_schema.sql"
FIXTURE_PATH = SCRIPT_DIR / "fixtures" / "M1" / "analyse.xlsx"

# Canonical file on the Synology share `math`, same folder the workbook
# records as D:\math\M1\. Compulsory core stays at math/HKDSE/Analyse.xls
# and is intentionally not opened here.
M1_RELATIVE = Path("math") / "M1" / "analyse.xlsx"
CORE_RELATIVE = Path("math") / "HKDSE" / "Analyse.xls"

YEAR_HEADER = re.compile(r"^(?:\d{2,4}|[A-Za-z]{1,6})$")
TOPIC_LIST = re.compile(r"^\d+(?:\s*,\s*\d+)*$")

# ms: M&S / older syllabus. dse: 2012 and after, including this workbook's SP/PP and ON.
SYLLABUS_MS = "ms"
SYLLABUS_DSE = "dse"
MS_TOPICS = (1, 16)
DSE_TOPICS = (1, 12)


@dataclass(frozen=True)
class LegendTopic:
    code: int
    short_name: str | None
    name: str | None


@dataclass(frozen=True)
class Unit:
    syllabus: str
    code: int
    short_name: str | None
    name: str | None


@dataclass(frozen=True)
class Question:
    sheet_name: str
    year_key: str
    year_label: str
    label_occurrence: int
    source_column: str
    column_index: int
    source_row: int
    question_no: int
    raw_cell: str
    unit_codes: tuple[int, ...]
    syllabus: str


@dataclass(frozen=True)
class Notice:
    level: str
    message: str


@dataclass(frozen=True)
class WorkbookMap:
    path: Path
    units: tuple[Unit, ...]
    questions: tuple[Question, ...]
    notices: tuple[Notice, ...]

    @property
    def errors(self) -> tuple[Notice, ...]:
        return tuple(notice for notice in self.notices if notice.level == "error")


def candidate_paths() -> list[Path]:
    """Search order after --xlsx and M1_ANALYSE_XLSX."""
    home = Path.home()
    return [
        Path("/volume1") / M1_RELATIVE,
        Path("/Volumes/math/M1/analyse.xlsx"),
        Path("/Volumes/DiskStation") / M1_RELATIVE,
        home / "SynologyDrive" / M1_RELATIVE,
        home / "DiskStation" / M1_RELATIVE,
        home / "Desktop" / M1_RELATIVE,
        Path("D:/math/M1/analyse.xlsx"),
        Path.cwd() / M1_RELATIVE,
        FIXTURE_PATH,
    ]


def resolve_xlsx(explicit: str | None = None, env: str | None = None) -> Path:
    if explicit:
        path = Path(explicit).expanduser()
        if not path.is_file():
            raise FileNotFoundError(f"M1 workbook not found: {path}")
        return path
    if env:
        path = Path(env).expanduser()
        if not path.is_file():
            raise FileNotFoundError(f"M1_ANALYSE_XLSX does not exist: {path}")
        return path
    for path in candidate_paths():
        if path.is_file():
            return path
    searched = "\n".join(f"  - {path}" for path in candidate_paths())
    raise FileNotFoundError(
        "M1 workbook math/M1/analyse.xlsx was not found. Pass --xlsx or set "
        f"M1_ANALYSE_XLSX.\nLooked in:\n{searched}"
    )


def calendar_year(label: str) -> int | None:
    """Map a header to a calendar year. 94 → 1994, 00 → 2000, 12 → 2012."""
    if not re.fullmatch(r"\d{2,4}", label):
        return None
    year = int(label)
    if year < 100:
        year += 1900 if year >= 90 else 2000
    return year


def syllabus_for(label: str, occurrence: int) -> str:
    """Topic-number range for one year column.

    The sheet has two columns labelled 12. Syllabus follows the column, not
    the shared label:

    - 94–99 and 00–11, and the first column labelled 12, are M&S.
      Allowed topics are 1–16. Topic 13 on that first 2012 column is valid.
    - The second column labelled 12, and calendar years after 2012, are DSE M1.
      Allowed topics are 1–12.
    - SP is the 2011 sample paper and PP is the 2022 practice paper.
      Both stay on topics 1–12.
    - ON is a 2012-and-after paper code. Allowed topics are 1–12.
    """
    token = label.strip().upper()
    if token == "ON" or token in {"SP", "PP"}:
        return SYLLABUS_DSE
    year = calendar_year(label.strip())
    if year is None:
        raise ValueError(f"year label {label!r} has no M1 syllabus rule")
    if year < 2012 or (year == 2012 and occurrence == 1):
        return SYLLABUS_MS
    return SYLLABUS_DSE


def column_sequence_key(label: str, occurrence: int, column_index: int) -> tuple[int, int, int, int]:
    """Order year columns as pre-2012, then SP, PP, then the two 2012 columns.

    The workbook lays SP and PP after the columns labelled 12. Seed order
    is SP (sample paper, 2011), PP (practice paper, 2022), then the first
    12 (M&S) and the second 12 (DSE M1), then 13 onward.
    """
    token = label.strip().upper()
    if token == "SP":
        return (1, 0, occurrence, column_index)
    if token == "PP":
        return (1, 1, occurrence, column_index)
    if token == "ON":
        return (1, 2, occurrence, column_index)
    year = calendar_year(label.strip())
    if year is not None and year < 2012:
        return (0, year, occurrence, column_index)
    if year is not None:
        return (2, year, occurrence, column_index)
    return (3, 0, occurrence, column_index)


def topic_bounds(syllabus: str) -> tuple[int, int]:
    if syllabus == SYLLABUS_MS:
        return MS_TOPICS
    if syllabus == SYLLABUS_DSE:
        return DSE_TOPICS
    raise ValueError(f"unknown syllabus {syllabus!r}")


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ValueError(f"unexpected boolean cell {value!r}")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"unexpected fractional cell {value!r}")
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def parse_topic_codes(value: object) -> tuple[int, ...]:
    text = cell_text(value)
    if not text:
        return ()
    if not TOPIC_LIST.fullmatch(text):
        raise ValueError(f"topic cell {text!r} is not a topic number from the workbook")
    return tuple(int(part) for part in re.split(r"\s*,\s*", text))


def is_year_label(value: object) -> bool:
    text = cell_text(value) if value is not None else ""
    return bool(text) and YEAR_HEADER.fullmatch(text) is not None


def is_name(value: object) -> bool:
    if value is None or isinstance(value, (int, float, bool)):
        return False
    text = str(value).strip()
    return bool(text) and TOPIC_LIST.fullmatch(text) is None


def question_number(value: object) -> int | None:
    if value is None:
        return None
    text = cell_text(value)
    if not text or not re.fullmatch(r"\d+", text):
        return None
    return int(text)


def load_m1_workbook(path: Path) -> WorkbookMap:
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Unknown extension is not supported and will be removed",
            category=UserWarning,
        )
        workbook = load_workbook(path, data_only=False, read_only=False)

    notices: list[Notice] = []
    legend_by_code: dict[int, LegendTopic] = {}
    questions: list[Question] = []

    try:
        for sheet in workbook.worksheets:
            parsed = _parse_sheet(sheet, notices)
            if parsed is None:
                continue
            sheet_legend, sheet_questions = parsed
            for topic in sheet_legend:
                existing = legend_by_code.get(topic.code)
                if existing is None:
                    legend_by_code[topic.code] = topic
                elif existing != topic:
                    raise ValueError(
                        f"topic {topic.code} has conflicting legend rows: "
                        f"{existing} vs {topic}"
                    )
            questions.extend(sheet_questions)
    finally:
        workbook.close()

    notices.extend(_range_errors(questions))
    for code in sorted(legend_by_code):
        if code < 1 or code > MS_TOPICS[1]:
            notices.append(
                Notice(
                    "error",
                    f"legend topic {code} is outside the older-syllabus range "
                    f"{MS_TOPICS[0]}-{MS_TOPICS[1]}",
                )
            )
    units = scoped_units(legend_by_code)
    return WorkbookMap(path=path, units=units, questions=tuple(questions), notices=tuple(notices))


def scoped_units(legend_by_code: dict[int, LegendTopic]) -> tuple[Unit, ...]:
    """Seed ms topics 1–16 and dse topics 1–12. Names come from the legend."""
    rows: list[Unit] = []
    for syllabus, (_low, high) in ((SYLLABUS_MS, MS_TOPICS), (SYLLABUS_DSE, DSE_TOPICS)):
        for code in range(1, high + 1):
            topic = legend_by_code.get(code)
            rows.append(
                Unit(
                    syllabus=syllabus,
                    code=code,
                    short_name=None if topic is None else topic.short_name,
                    name=None if topic is None else topic.name,
                )
            )
    return tuple(rows)


def _range_errors(questions: list[Question]) -> list[Notice]:
    notices: list[Notice] = []
    for question in questions:
        low, high = topic_bounds(question.syllabus)
        outside = [str(code) for code in question.unit_codes if code < low or code > high]
        if not outside:
            continue
        notices.append(
            Notice(
                "error",
                f"{question.sheet_name}!{question.source_column}{question.source_row} "
                f"{question.year_key} Q{question.question_no} syllabus {question.syllabus}: "
                f"topic {', '.join(outside)} is outside {low}-{high}",
            )
        )
    return notices


def _parse_sheet(sheet, notices: list[Notice]) -> tuple[list[LegendTopic], list[Question]] | None:
    header_row = _find_header_row(sheet)
    if header_row is None:
        notices.append(Notice("info", f"{sheet.title}: no year header, skipped"))
        return None

    year_columns: list[tuple[int, str, str]] = []
    for column in range(2, sheet.max_column + 1):
        value = sheet.cell(header_row, column).value
        if value is None or cell_text(value) == "":
            continue
        if not is_year_label(value):
            raise ValueError(
                f"{sheet.title}!{get_column_letter(column)}{header_row} "
                f"header {value!r} is not a year or paper code"
            )
        year_columns.append((column, get_column_letter(column), cell_text(value)))

    if not year_columns:
        notices.append(Notice("info", f"{sheet.title}: no year columns, skipped"))
        return None

    legend_start = _find_legend_start(sheet, header_row, year_columns)
    if legend_start is None:
        last_question = sheet.max_row
    else:
        last_question = legend_start - 1

    questions = _parse_questions(sheet, header_row, last_question, year_columns)
    units = _parse_legend(sheet, legend_start) if legend_start is not None else []
    if not questions and not units:
        notices.append(Notice("info", f"{sheet.title}: empty grid, skipped"))
        return None

    notices.append(
        Notice(
            "info",
            f"{sheet.title}: {len(questions)} questions, {len(year_columns)} year columns, "
            f"{len(units)} legend topics",
        )
    )
    return units, questions


def _find_header_row(sheet) -> int | None:
    limit = min(sheet.max_row, 5)
    for row in range(1, limit + 1):
        labels = [
            cell_text(sheet.cell(row, column).value)
            for column in range(2, sheet.max_column + 1)
            if sheet.cell(row, column).value is not None
        ]
        if sum(1 for label in labels if YEAR_HEADER.fullmatch(label)) >= 3:
            return row
    return None


def _find_legend_start(sheet, header_row: int, year_columns: list[tuple[int, str, str]]) -> int | None:
    for row in range(header_row + 1, sheet.max_row + 1):
        if question_number(sheet.cell(row, 1).value) is None:
            continue
        if is_name(sheet.cell(row, 2).value) or is_name(sheet.cell(row, 6).value):
            return row
        # A later block of bare topic numbers (no names yet) still belongs to
        # the legend when the year columns on that row are empty.
        if all(sheet.cell(row, column).value is None for column, _, _ in year_columns):
            previous_is_question = False
            if row > header_row + 1 and question_number(sheet.cell(row - 1, 1).value) is not None:
                previous_is_question = any(
                    sheet.cell(row - 1, column).value is not None for column, _, _ in year_columns
                )
            if not previous_is_question:
                return row
    return None


def _parse_questions(
    sheet,
    header_row: int,
    last_question: int,
    year_columns: list[tuple[int, str, str]],
) -> list[Question]:
    labels = [label for _, _, label in year_columns]
    occurrences = Counter(labels)
    questions: list[Question] = []
    ordered_columns = sorted(
        year_columns,
        key=lambda item: column_sequence_key(
            item[2],
            _label_occurrence(year_columns, item[0]),
            item[0],
        ),
    )

    for row in range(header_row + 1, last_question + 1):
        number = question_number(sheet.cell(row, 1).value)
        if number is None:
            if any(sheet.cell(row, column).value is not None for column, _, _ in year_columns):
                raise ValueError(f"{sheet.title}!A{row} is blank but year cells are filled")
            continue
        for column, letter, label in ordered_columns:
            value = sheet.cell(row, column).value
            if value is None or cell_text(value) == "":
                continue
            try:
                codes = parse_topic_codes(value)
            except ValueError as exc:
                raise ValueError(f"{sheet.title}!{letter}{row}: {exc}") from exc
            occurrence = _label_occurrence(year_columns, column)
            year_key = label if occurrences[label] == 1 else f"{label}#{letter}"
            questions.append(
                Question(
                    sheet_name=sheet.title,
                    year_key=year_key,
                    year_label=label,
                    label_occurrence=occurrence,
                    source_column=letter,
                    column_index=column,
                    source_row=row,
                    question_no=number,
                    raw_cell=cell_text(value),
                    unit_codes=codes,
                    syllabus=syllabus_for(label, occurrence),
                )
            )
    return questions


def _label_occurrence(year_columns: list[tuple[int, str, str]], column: int) -> int:
    label = next(item[2] for item in year_columns if item[0] == column)
    occurrence = 0
    for candidate, _, candidate_label in year_columns:
        if candidate_label == label:
            occurrence += 1
        if candidate == column:
            return occurrence
    raise AssertionError(f"column {column} missing from year headers")


def _parse_legend(sheet, legend_start: int) -> list[LegendTopic]:
    topics: list[LegendTopic] = []
    seen: set[int] = set()
    blank_run = 0
    for row in range(legend_start, sheet.max_row + 1):
        number = question_number(sheet.cell(row, 1).value)
        if number is None:
            blank_run += 1
            if blank_run >= 1 and topics:
                # One blank row ends the legend block.
                if all(sheet.cell(row, column).value is None for column in range(1, 7)):
                    break
            continue
        blank_run = 0
        if number in seen:
            raise ValueError(f"{sheet.title}!A{row} repeats legend topic number {number}")
        seen.add(number)
        short_name = _optional_name(sheet.cell(row, 2).value)
        name = _optional_name(sheet.cell(row, 6).value)
        topics.append(LegendTopic(number, short_name, name))
    return topics


def _optional_name(value: object) -> str | None:
    if not is_name(value):
        return None
    text = str(value).strip()
    return text or None


def render_sql(mapping: WorkbookMap) -> str:
    if mapping.errors:
        details = "\n".join(notice.message for notice in mapping.errors)
        raise ValueError(f"refusing to render SQL with parse errors:\n{details}")

    lines = [
        "-- Generated by db/seed_analyse_m1.py",
        f"-- Source: {mapping.path}",
        "-- Topic numbers are the workbook codes, scoped by syllabus.",
        "-- ms topics are 1-16. dse topics are 1-12. Compulsory analyse_* tables are not touched.",
        "SET NAMES utf8mb4;",
        "START TRANSACTION;",
        "DELETE FROM analyse_m1_sq_units;",
        "DELETE FROM analyse_m1_sq;",
        "DELETE FROM analyse_m1_units;",
    ]
    if mapping.units:
        lines.append(
            "INSERT INTO analyse_m1_units (syllabus, unit_code, short_name, name) VALUES"
        )
        lines.append(
            ",\n".join(_unit_values(unit) for unit in mapping.units) + ";"
        )
    if mapping.questions:
        lines.append(
            "INSERT INTO analyse_m1_sq "
            "(id, sheet_name, year_key, year_label, label_occurrence, source_column, "
            "question_no, raw_cell, syllabus) VALUES"
        )
        question_rows = []
        for sq_id, question in enumerate(mapping.questions, start=1):
            question_rows.append(
                "("
                f"{sq_id}, {sql_str(question.sheet_name)}, {sql_str(question.year_key)}, "
                f"{sql_str(question.year_label)}, {question.label_occurrence}, "
                f"{sql_str(question.source_column)}, {question.question_no}, "
                f"{sql_str(question.raw_cell)}, {sql_str(question.syllabus)}"
                ")"
            )
        lines.append(",\n".join(question_rows) + ";")
        lines.append(
            "INSERT INTO analyse_m1_sq_units (sq_id, position, syllabus, unit_code) VALUES"
        )
        link_rows = []
        for sq_id, question in enumerate(mapping.questions, start=1):
            for position, code in enumerate(question.unit_codes, start=1):
                link_rows.append(
                    f"({sq_id}, {position}, {sql_str(question.syllabus)}, {code})"
                )
        lines.append(",\n".join(link_rows) + ";")
    lines.append("COMMIT;")
    lines.append("")
    return "\n".join(lines)


def _unit_values(unit: Unit) -> str:
    return (
        f"({sql_str(unit.syllabus)}, {unit.code}, "
        f"{sql_nullable(unit.short_name)}, {sql_nullable(unit.name)})"
    )


def sql_nullable(value: str | None) -> str:
    if value is None:
        return "NULL"
    return sql_str(value)


def sql_str(value: str) -> str:
    escaped = value.replace("\\", "\\\\").replace("'", "''")
    return f"'{escaped}'"


def render_report(mapping: WorkbookMap) -> str:
    link_count = sum(len(question.unit_codes) for question in mapping.questions)
    year_keys = []
    seen: set[str] = set()
    for question in mapping.questions:
        if question.year_key not in seen:
            seen.add(question.year_key)
            year_keys.append(question.year_key)
    lines = [f"source: {mapping.path}"]
    for syllabus, bounds in ((SYLLABUS_MS, MS_TOPICS), (SYLLABUS_DSE, DSE_TOPICS)):
        group = [question for question in mapping.questions if question.syllabus == syllabus]
        group_links = sum(len(question.unit_codes) for question in group)
        group_units = [unit.code for unit in mapping.units if unit.syllabus == syllabus]
        lines.append(
            f"syllabus {syllabus} topics {bounds[0]}-{bounds[1]}: "
            f"units {len(group_units)} ({group_units[0]}-{group_units[-1]}), "
            f"questions {len(group)}, topic links {group_links}"
        )
    lines.extend(
        [
            f"questions: {len(mapping.questions)}",
            f"topic links: {link_count}",
            "year_key: " + ", ".join(year_keys),
        ]
    )
    for notice in mapping.notices:
        lines.append(f"{notice.level}: {notice.message}")
    if mapping.errors:
        lines.append("result: failed")
    else:
        lines.append("result: ok")
    return "\n".join(lines)


def apply_sql(sql: str) -> None:
    mysql = os.environ.get("MYSQL_BIN", "mysql")
    database = os.environ.get("MYSQL_DATABASE", "")
    if not database:
        raise SystemExit(
            "Refusing to apply: set MYSQL_DATABASE (and MYSQL_HOST, MYSQL_PORT, "
            "MYSQL_USER, MYSQL_PWD) on the machine that hosts alphaeqt. "
            "This script does not contact a database during --dry-run."
        )
    command = [
        mysql,
        "--batch",
        "--connect-timeout=10",
        "-h",
        os.environ.get("MYSQL_HOST", "127.0.0.1"),
        "-P",
        os.environ.get("MYSQL_PORT", "3306"),
        "-u",
        os.environ.get("MYSQL_USER", "root"),
    ]
    defaults = os.environ.get("MYSQL_DEFAULTS_FILE")
    if defaults:
        command.extend([f"--defaults-extra-file={defaults}"])
    command.append(database)
    completed = subprocess.run(
        command,
        input=sql,
        text=True,
        check=False,
    )
    if completed.returncode != 0:
        raise SystemExit(f"mysql exited {completed.returncode}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Import M1 Analyse year x question topic numbers from "
            "math/M1/analyse.xlsx into analyse_m1_* tables. "
            "The default is a dry run and does not touch MySQL. "
            f"Compulsory core workbook {CORE_RELATIVE.as_posix()} is not read."
        )
    )
    parser.add_argument(
        "--xlsx",
        help=(
            "Path to M1 analyse.xlsx. Overrides M1_ANALYSE_XLSX and the "
            "Synology search path. Canonical location: math/M1/analyse.xlsx "
            "on the DiskStation share (often /Volumes/math/M1/analyse.xlsx "
            "or D:\\math\\M1\\analyse.xlsx)."
        ),
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Parse and print a report. This is the default.",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help=(
            "Parse and exit 1 if a cell is not a topic number in range for that "
            "column: 1-16 before 2012 and on the first column labelled 12; "
            "1-12 on the second 12, later years, SP, PP, and ON."
        ),
    )
    parser.add_argument(
        "--emit-sql",
        action="store_true",
        help="Print SQL that replaces analyse_m1_* rows. Does not connect.",
    )
    parser.add_argument(
        "--with-schema",
        action="store_true",
        help="With --emit-sql or --apply, prepend db/analyse_m1_schema.sql.",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help=(
            "Apply schema and seed SQL through the local mysql client. "
            "Requires MYSQL_DATABASE. Do not point this at production from a "
            "machine that does not already have the alphaeqt credentials."
        ),
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    selected = sum(bool(flag) for flag in (args.check, args.emit_sql, args.apply))
    if args.dry_run and selected:
        print("--dry-run cannot be combined with --check, --emit-sql, or --apply", file=sys.stderr)
        return 2
    if selected > 1:
        print("use only one of --check, --emit-sql, --apply", file=sys.stderr)
        return 2

    try:
        path = resolve_xlsx(args.xlsx, os.environ.get("M1_ANALYSE_XLSX"))
        mapping = load_m1_workbook(path)
    except (FileNotFoundError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    if mapping.errors:
        print(render_report(mapping), file=sys.stderr)
        return 1

    if args.emit_sql or args.apply:
        sql = render_sql(mapping)
        if args.with_schema or args.apply:
            sql = SCHEMA_PATH.read_text(encoding="utf-8") + "\n" + sql
        if args.emit_sql:
            sys.stdout.write(sql)
            return 0
        apply_sql(sql)
        print(render_report(mapping))
        return 0

    print(render_report(mapping))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
